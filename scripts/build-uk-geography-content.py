"""Compile pinned ONS evidence into immutable UK Geography Content v1."""

from __future__ import annotations

import hashlib
import json
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
CONTENT = ROOT / "src/data/geography/uk/primary-local-admin-2024"
ARTIFACTS = CONTENT / "artifacts"
PARTITION_ID = "geography.uk.primary-local-admin-2024-06-30-v1"
GEOGRAPHY_SOURCE_ID = "geography.source.ons.lad-may-2024-v1"
ADMIN_SOURCE_ID = "geography.source.ons.uk-administrative-geography-v1"
POPULATION_SOURCE_ID = "geography.source.ons.myeb-local-authorities-mid-2024-v1"
EXPECTED = {
    "ons-lad-may-2024-attributes.json": (29265, "a0dc2a30f66a87e9c12156ca37cbba6e726d8fa3ff166066381717d50954414c"),
    "ons-lad-may-2024-item.json": (4862, "8bc3a4e2ec1bc43f86e9afd3dfc8aae500b16db4e39130847e8a127310e52908"),
    "ons-lad-may-2024-layer.json": (13451, "e2ff820533c7bd4b168cf015a5c244b4c44ace5ee889d60c78a11161487238c9"),
    "ons-myeb-local-authorities-mid-2024.xlsx": (47036552, "321f27261c9cf110ea44ca10d198580ec3969bc1171c780cba5596184df9b04d"),
}
COUNTRIES = {
    "E": ("england", "place.uk.constituent.england"),
    "W": ("wales", "place.uk.constituent.wales"),
    "S": ("scotland", "place.uk.constituent.scotland"),
    "N": ("northern-ireland", "place.uk.constituent.northern-ireland"),
}
EXPECTED_KIND_COUNTS = Counter({
    "geography.uk.england.unitary-authority": 62,
    "geography.uk.england.non-metropolitan-district": 164,
    "geography.uk.england.metropolitan-district": 36,
    "geography.uk.england.london-borough": 32,
    "geography.uk.england.city-of-london": 1,
    "geography.uk.england.isles-of-scilly": 1,
    "geography.uk.wales.principal-area": 22,
    "geography.uk.scotland.council-area": 32,
    "geography.uk.northern-ireland.local-government-district": 11,
})
ROOTS = [
    ("place.uk.country.united-kingdom", "United Kingdom", "geography.uk.root", None),
    ("place.uk.constituent.england", "England", "geography.uk.constituent-country", "place.uk.country.united-kingdom"),
    ("place.uk.constituent.wales", "Wales", "geography.uk.constituent-country", "place.uk.country.united-kingdom"),
    ("place.uk.constituent.scotland", "Scotland", "geography.uk.constituent-country", "place.uk.country.united-kingdom"),
    ("place.uk.constituent.northern-ireland", "Northern Ireland", "geography.uk.constituent-country", "place.uk.country.united-kingdom"),
]


def encode(value: object) -> bytes:
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"), sort_keys=True).encode("utf-8")


def write_json(name: str, value: object) -> None:
    (CONTENT / name).write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def fnv1a64(value: object) -> str:
    result = 0xCBF29CE484222325
    for byte in encode(value):
        result ^= byte
        result = (result * 0x100000001B3) & 0xFFFFFFFFFFFFFFFF
    return f"fnv1a64-v1:{result:016x}"


def normalized_text(value: object, field: str) -> str:
    if not isinstance(value, str):
        raise ValueError(f"{field} is not text")
    result = unicodedata.normalize("NFC", value.strip())
    if not result:
        raise ValueError(f"{field} is empty")
    return result


def kind(code: str) -> str:
    if code == "E09000001":
        return "geography.uk.england.city-of-london"
    if code == "E06000053":
        return "geography.uk.england.isles-of-scilly"
    return {
        "E06": "geography.uk.england.unitary-authority",
        "E07": "geography.uk.england.non-metropolitan-district",
        "E08": "geography.uk.england.metropolitan-district",
        "E09": "geography.uk.england.london-borough",
        "W06": "geography.uk.wales.principal-area",
        "S12": "geography.uk.scotland.council-area",
        "N09": "geography.uk.northern-ireland.local-government-district",
    }.get(code[:3], "")


def verify_artifacts() -> None:
    for filename, (size, digest) in EXPECTED.items():
        path = ARTIFACTS / filename
        if not path.is_file() or path.stat().st_size != size or sha256(path) != digest:
            raise ValueError(f"Pinned artifact mismatch: {filename}")


def load_geography_rows() -> list[dict[str, object]]:
    item = json.loads((ARTIFACTS / "ons-lad-may-2024-item.json").read_text(encoding="utf-8"))
    if (
        item.get("id") != "f3528c2d6d454edab74f2648cc6a45f6"
        or item.get("owner") != "ONSGeography_data"
        or item.get("title") != "Local Authority Districts (May 2024) Boundaries UK BGC"
        or item.get("type") != "Feature Service"
        or item.get("contentStatus") != "public_authoritative"
    ):
        raise ValueError("ONS item metadata changed")
    layer = json.loads((ARTIFACTS / "ons-lad-may-2024-layer.json").read_text(encoding="utf-8"))
    expected_fields = [
        "FID", "LAD24CD", "LAD24NM", "LAD24NMW", "BNG_E", "BNG_N",
        "LONG", "LAT", "Shape__Area", "Shape__Length", "GlobalID",
    ]
    if (
        layer.get("name") != "LAD_MAY_2024_UK_BGC"
        or layer.get("type") != "Feature Layer"
        or layer.get("serviceItemId") != item["id"]
        or layer.get("editingInfo", {}).get("dataLastEditDate") != 1719925309918
        or [field.get("name") for field in layer.get("fields", [])] != expected_fields
    ):
        raise ValueError("ONS layer schema or vintage changed")
    source = json.loads((ARTIFACTS / "ons-lad-may-2024-attributes.json").read_text(encoding="utf-8"))
    if set(source) != {
        "objectIdFieldName", "uniqueIdField", "globalIdFieldName", "serverGens",
        "geometryType", "spatialReference", "fields", "features",
    }:
        raise ValueError("ONS attribute artifact shape changed")
    if [field.get("name") for field in source.get("fields", [])] != ["LAD24CD", "LAD24NM", "LAD24NMW"]:
        raise ValueError("ONS attribute fields changed")
    features = source.get("features")
    if not isinstance(features, list) or len(features) != 361:
        raise ValueError("ONS attribute inventory is not 361 rows")
    records: list[dict[str, object]] = []
    seen: set[str] = set()
    for feature in features:
        if not isinstance(feature, dict) or set(feature) != {"attributes"}:
            raise ValueError("Malformed ONS feature row")
        attrs = feature["attributes"]
        if not isinstance(attrs, dict) or set(attrs) != {"LAD24CD", "LAD24NM", "LAD24NMW"}:
            raise ValueError("Malformed ONS feature attributes")
        code = normalized_text(attrs["LAD24CD"], "LAD24CD")
        if not __import__("re").fullmatch(r"[EWSN][0-9]{8}", code) or code in seen:
            raise ValueError(f"Invalid or duplicate GSS code: {code}")
        seen.add(code)
        country = COUNTRIES.get(code[0])
        unit_kind = kind(code)
        if country is None or not unit_kind:
            raise ValueError(f"Unsupported GSS unit type: {code}")
        welsh_raw = attrs["LAD24NMW"]
        if not isinstance(welsh_raw, str):
            raise ValueError("LAD24NMW is not text")
        welsh = unicodedata.normalize("NFC", welsh_raw.strip()) or None
        records.append({
            "officialCode": code,
            "displayName": normalized_text(attrs["LAD24NM"], "LAD24NM"),
            "welshDisplayName": welsh,
            "countryBranch": country[0],
            "kindId": unit_kind,
            "placeId": f"place.uk.local-admin.{code.lower()}",
            "parentPlaceId": country[1],
        })
    records.sort(key=lambda row: str(row["officialCode"]))
    country_counts = Counter(str(row["countryBranch"]) for row in records)
    if country_counts != Counter({"england": 296, "wales": 22, "scotland": 32, "northern-ireland": 11}):
        raise ValueError(f"Unexpected country inventory: {country_counts}")
    kind_counts = Counter(str(row["kindId"]) for row in records)
    if kind_counts != EXPECTED_KIND_COUNTS:
        raise ValueError(f"Unexpected source-kind inventory: {kind_counts}")
    return records


def population_index(records: list[dict[str, object]]) -> list[dict[str, str]]:
    workbook = openpyxl.load_workbook(
        ARTIFACTS / "ons-myeb-local-authorities-mid-2024.xlsx",
        read_only=True,
        data_only=True,
    )
    if workbook.sheetnames != ["Cover sheet", "Contents", "Notes", "Related publications", "MYEB1", "MYEB2", "MYEB3", "MYEB4", "MYEB5"]:
        raise ValueError("Population workbook sheet inventory changed")
    sheet = workbook["MYEB1"]
    header = [sheet.cell(2, column).value for column in range(1, 20)]
    if header[:5] != ["ladcode23", "laname23", "country", "sex", "age"] or header[-1] != "population_2024":
        raise ValueError("Population workbook MYEB1 schema changed")
    index: dict[str, tuple[str, str]] = {}
    row_counts: Counter[str] = Counter()
    for row in sheet.iter_rows(min_row=3, values_only=True):
        code = row[0]
        if code is None:
            continue
        if not isinstance(code, str) or not __import__("re").fullmatch(r"[EWSN][0-9]{8}", code):
            raise ValueError("Population workbook contains a malformed area code")
        name = normalized_text(row[1], "laname23")
        country = normalized_text(row[2], "country")
        previous = index.setdefault(code, (name, country))
        if previous != (name, country):
            raise ValueError(f"Population workbook area identity changes within rows: {code}")
        row_counts[code] += 1
    if len(index) != 361 or any(count != 182 for count in row_counts.values()):
        raise ValueError("Population workbook local-area inventory changed")
    geography = {str(row["officialCode"]): row for row in records}
    if set(index) != set(geography):
        raise ValueError("Geography and population code sets differ")
    output: list[dict[str, str]] = []
    for code in sorted(index):
        name, country = index[code]
        source = geography[code]
        if name != source["displayName"] or country != code[0]:
            raise ValueError(f"Geography/population identity mismatch: {code}")
        output.append({"officialCode": code, "displayName": name, "countryCode": country})
    return output


def main() -> None:
    verify_artifacts()
    records = load_geography_rows()
    population = population_index(records)
    artifact_meta = {
        "ons-lad-may-2024-item.json": (
            "artifact.ons.lad-may-2024-item-v1", "Official ArcGIS item metadata for the authoritative May 2024 feature service.",
            "https://www.arcgis.com/sharing/rest/content/items/f3528c2d6d454edab74f2648cc6a45f6?f=json", "application/json",
        ),
        "ons-lad-may-2024-layer.json": (
            "artifact.ons.lad-may-2024-layer-v1", "Official layer schema and immutable source-vintage markers.",
            "https://services1.arcgis.com/ESMARspQHYMw9BZ9/arcgis/rest/services/Local_Authority_Districts_May_2024_Boundaries_UK_BGC/FeatureServer/0?f=pjson", "application/json",
        ),
        "ons-lad-may-2024-attributes.json": (
            "artifact.ons.lad-may-2024-attributes-v1", "Official code, primary-name and Welsh alternate-name attributes; geometry deliberately excluded.",
            "https://services1.arcgis.com/ESMARspQHYMw9BZ9/arcgis/rest/services/Local_Authority_Districts_May_2024_Boundaries_UK_BGC/FeatureServer/0/query?f=json&where=1%3D1&outFields=LAD24CD%2CLAD24NM%2CLAD24NMW&returnGeometry=false&orderByFields=LAD24CD+ASC", "application/json",
        ),
        "ons-myeb-local-authorities-mid-2024.xlsx": (
            "artifact.ons.myeb-local-authorities-mid-2024-v1", "Integrity cross-check of the 361 area codes, names and country assignments; population counts do not enter Geography.",
            "https://www.ons.gov.uk/file?uri=%2Fpeoplepopulationandcommunity%2Fpopulationandmigration%2Fpopulationestimates%2Fdatasets%2Fpopulationestimatesforukenglandandwalesscotlandandnorthernireland%2Fmid2011tomid2024%2Fmyebtablesuk20112024.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        ),
    }
    source_manifest = {
        "version": 1,
        "bundleId": "evidence.uk-geography.primary-local-admin-2024-v1",
        "retrievalDate": "2026-09-28",
        "licence": "Open Government Licence v3.0",
        "licenceLocator": "https://www.nationalarchives.gov.uk/doc/open-government-licence/version/3/",
        "attribution": "Source: Office for National Statistics licensed under the Open Government Licence v3.0. Contains OS data © Crown copyright and database right 2024.",
        "artifacts": [
            {
                "id": artifact_meta[filename][0],
                "producer": "Office for National Statistics",
                "product": "Local Authority Districts (May 2024) Boundaries UK BGC" if "lad-may" in filename else "Population estimates for the UK, England, Wales, Scotland and Northern Ireland: mid-2024",
                "releaseId": "ons-open-geography-lad-may-2024" if "lad-may" in filename else "MYE24UK-detailed",
                "referenceDate": "2024-05-01" if "lad-may" in filename else "2024-06-30",
                "filename": filename,
                "mediaType": artifact_meta[filename][3],
                "size": EXPECTED[filename][0],
                "sha256": EXPECTED[filename][1],
                "sourceLocator": artifact_meta[filename][2],
                "purpose": artifact_meta[filename][1],
                "expectedSchema": "ArcGIS item metadata" if filename.endswith("item.json") else "ArcGIS feature-layer metadata" if filename.endswith("layer.json") else "ArcGIS query with LAD24CD, LAD24NM and LAD24NMW" if filename.endswith("attributes.json") else "XLSX sheets Cover sheet, Contents, Notes, Related publications, MYEB1..MYEB5; MYEB1 local-authority identity columns",
            }
            for filename in sorted(EXPECTED)
        ],
    }
    write_json("source-manifest.json", source_manifest)
    normalized = {
        "version": 1,
        "sourceArtifactId": "artifact.ons.lad-may-2024-attributes-v1",
        "normalization": "Unicode NFC and source-format surrounding-whitespace removal only; punctuation, capitalization and diacritics preserved; no transliteration or fuzzy correction.",
        "records": records,
    }
    write_json("normalized-local-authorities.json", normalized)
    population_output = {
        "version": 1,
        "sourceArtifactId": "artifact.ons.myeb-local-authorities-mid-2024-v1",
        "workbookSha256": EXPECTED["ons-myeb-local-authorities-mid-2024.xlsx"][1],
        "sheet": "MYEB1",
        "records": population,
    }
    write_json("population-area-index.json", population_output)
    continuity_entries = [
        {
            "sourceArtifactId": "artifact.ons.lad-may-2024-attributes-v1",
            "sourceCode": str(row["officialCode"]),
            "placeId": str(row["placeId"]),
            "decision": "new-place",
            "reviewBasis": "Initial accepted May 2024 official identity; future continuity requires a separate reviewed decision.",
        }
        for row in records
    ]
    continuity_semantics = {"version": 1, "entries": continuity_entries}
    continuity = {
        "version": 1,
        "id": "geography.uk.place-continuity-2024-v1",
        "fingerprint": fnv1a64(continuity_semantics),
        "entries": continuity_entries,
    }
    write_json("continuity-ledger.json", continuity)
    places = [{"placeId": place_id, "countryId": "uk"} for place_id, _, _, _ in ROOTS]
    places.extend({"placeId": str(row["placeId"]), "countryId": "uk"} for row in records)
    places.sort(key=lambda row: row["placeId"])
    write_json("place-identities.json", {"version": 1, "places": places})
    place_manifest = [
        {
            "placeId": place["placeId"],
            "fingerprint": fnv1a64({"version": 1, "placeId": place["placeId"], "countryId": "uk"}),
        }
        for place in places
    ]
    place_manifest_semantics = {"version": 1, "entries": place_manifest}
    place_manifest_output = {
        "version": 1,
        "fingerprint": fnv1a64(place_manifest_semantics),
        "entries": place_manifest,
    }
    write_json("place-manifest.json", place_manifest_output)
    sources = [
        {
            "version": 1,
            "id": ADMIN_SOURCE_ID,
            "producer": "Office for National Statistics",
            "datasetId": "ons.uk-administrative-geography",
            "releaseId": "administrative-geography-reviewed-2024",
            "title": "Administrative geography of the United Kingdom",
            "jurisdiction": "uk",
            "referenceDate": {"year": 2024, "month": 6, "day": 30},
            "classification": "statutory-institutional",
            "methodology": "Official structural reference for the UK root, constituent countries and differing local-government forms. It is not a claim that the 361 allocation units share one legal type.",
            "licence": "Open Government Licence v3.0",
            "externalLocator": "https://www.ons.gov.uk/methodology/geography/ukgeographies/administrativegeography",
        },
        {
            "version": 1,
            "id": GEOGRAPHY_SOURCE_ID,
            "producer": "Office for National Statistics",
            "datasetId": "ons-open-geography.local-authority-districts-may-2024-uk-bgc",
            "releaseId": "ons-open-geography-lad-may-2024",
            "title": "Local Authority Districts (May 2024) Boundaries UK BGC",
            "jurisdiction": "uk",
            "referenceDate": {"year": 2024, "month": 5, "day": 1},
            "classification": "statutory-institutional",
            "methodology": "Official May 2024 Local Authority District code and name inventory. Only code/name attributes are compiled; boundary geometry and coordinates are excluded.",
            "licence": source_manifest["licence"] + ". " + source_manifest["attribution"],
            "bundledArtifact": {"id": "artifact.ons.lad-may-2024-attributes-v1", "sha256": "sha256:" + EXPECTED["ons-lad-may-2024-attributes.json"][1]},
            "externalLocator": artifact_meta["ons-lad-may-2024-attributes.json"][2],
        },
        {
            "version": 1,
            "id": POPULATION_SOURCE_ID,
            "producer": "Office for National Statistics",
            "datasetId": "ons.population-estimates-uk-local-authorities-mid-2024",
            "releaseId": "MYE24UK-detailed",
            "title": "Annual mid-year population estimates for United Kingdom local authorities, 2011 to 2024",
            "jurisdiction": "uk",
            "referenceDate": {"year": 2024, "month": 6, "day": 30},
            "classification": "estimated",
            "methodology": "Used only to prove exact equality of the 361 official code/name/country identities with the frozen mid-2024 population evidence. Population counts do not enter Geography content.",
            "licence": "Open Government Licence v3.0",
            "bundledArtifact": {"id": "artifact.ons.myeb-local-authorities-mid-2024-v1", "sha256": "sha256:" + EXPECTED["ons-myeb-local-authorities-mid-2024.xlsx"][1]},
            "externalLocator": artifact_meta["ons-myeb-local-authorities-mid-2024.xlsx"][2],
        },
    ]
    sources.sort(key=lambda row: row["id"])
    nodes = [
        {
            "placeId": place_id,
            "displayName": display_name,
            "kindId": kind_id,
            "parentPlaceId": parent,
            "populationAllocationCell": False,
            "sourceIds": [ADMIN_SOURCE_ID],
        }
        for place_id, display_name, kind_id, parent in ROOTS
    ]
    nodes.extend(
        {
            "placeId": row["placeId"],
            "displayName": row["displayName"],
            "kindId": row["kindId"],
            "parentPlaceId": row["parentPlaceId"],
            "populationAllocationCell": True,
            "sourceIds": sorted([GEOGRAPHY_SOURCE_ID, POPULATION_SOURCE_ID]),
        }
        for row in records
    )
    nodes.sort(key=lambda row: str(row["placeId"]))
    limitations = sorted([
        "Administrative allocation geography only; nodes do not assert settlement, address, residence or neighbourhood meaning.",
        "Boundary geometry, coordinates, adjacency, distance, routes and travel semantics are absent.",
        "The 361 allocation units play one functional Population role but retain different official legal and administrative kinds.",
        "The official area inventory is the May 2024 ONS Geography snapshot and is applied to the aligned 30 June 2024 population reference date.",
        "English regions, upper-tier counties, settlements, wards, constituencies, postcodes and statistical small areas are excluded.",
        "Jersey, Guernsey, the Isle of Man and UK Overseas Territories are outside this UK demographic universe.",
        "Population remains geographically unallocated; this package contains no demographic counts and changes no Population state.",
    ])
    partition_without_fingerprint = {
        "version": 1,
        "partitionId": PARTITION_ID,
        "countryId": "uk",
        "effectiveDate": {"year": 2024, "month": 6, "day": 30},
        "sources": sources,
        "nodes": nodes,
        "limitations": limitations,
    }
    semantic = {key: value for key, value in partition_without_fingerprint.items() if key != "partitionId"}
    partition = dict(partition_without_fingerprint)
    partition["fingerprint"] = fnv1a64(semantic)
    partition = {key: partition[key] for key in ["version", "partitionId", "fingerprint", "countryId", "effectiveDate", "sources", "nodes", "limitations"]}
    write_json("compiled-partition.json", partition)
    package_manifest = {PARTITION_ID: partition["fingerprint"]}
    write_json("package-manifest.json", package_manifest)
    kind_counts = dict(sorted(Counter(str(row["kindId"]) for row in records).items()))
    country_counts = dict(sorted(Counter(str(row["countryBranch"]) for row in records).items()))
    serialized_size = len(encode(partition))
    report = {
        "version": 1,
        "packageId": PARTITION_ID,
        "effectiveDate": "2024-06-30",
        "sourceArtifacts": [{"filename": name, "size": value[0], "sha256": value[1]} for name, value in sorted(EXPECTED.items())],
        "placeCount": len(places),
        "nodeCount": len(nodes),
        "allocationCellCount": len(records),
        "countByCountry": country_counts,
        "countByKindId": kind_counts,
        "missingCodes": [],
        "unexpectedCodes": [],
        "duplicateCodes": [],
        "populationWorkbookComparison": {"codeSetEqual": True, "nameMismatchCount": 0, "countryMismatchCount": 0, "areaCount": len(population)},
        "continuityLedgerFingerprint": continuity["fingerprint"],
        "placeManifestFingerprint": place_manifest_output["fingerprint"],
        "partitionFingerprint": partition["fingerprint"],
        "packageSerializedSize": serialized_size,
    }
    write_json("build-report.json", report)


if __name__ == "__main__":
    main()
