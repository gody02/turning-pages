"""Build the unregistered UK mid-2024 geographic Population candidate."""
from __future__ import annotations

import hashlib, heapq, json, sys, time
from collections import defaultdict
from pathlib import Path
from fractions import Fraction
import openpyxl

ROOT=Path(__file__).resolve().parents[1]
GEO=ROOT/'src/data/geography/uk/primary-local-admin-2024'
OUT=ROOT/'src/data/demography/uk/geographic-mid-2024'
BOOK=GEO/'artifacts/ons-myeb-local-authorities-mid-2024.xlsx'
V2=ROOT/'src/data/demography/uk/ons-mid-2024/compiled-package-v2.json'
PACKAGE_ID='uk.population.mid-2024.v3'
PARTITION_ID='geography.uk.primary-local-admin-2024-06-30-v1'
PROFILE_ID='human.uk.pending-mid-2024-v1'
BOOK_HASH='321f27261c9cf110ea44ca10d198580ec3969bc1171c780cba5596184df9b04d'
TOTAL=69_281_437

def canonical(v): return json.dumps(v,ensure_ascii=False,sort_keys=True,separators=(',',':'))
def fnv(v):
    h=0xcbf29ce484222325
    for b in canonical(v).encode('utf-8'): h=((h^b)*0x100000001b3)&0xffffffffffffffff
    return f'fnv1a64-v1:{h:016x}'
def cohort_id(area,year):
    enc=lambda s:f'{len(s)}:{s}'
    return f'population-cohort:v1|{enc("uk")}|{enc(area)}|{year:04d}|{enc(PROFILE_ID)}'
def write(name,v):
    OUT.mkdir(parents=True,exist_ok=True); (OUT/name).write_text(json.dumps(v,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')

class Edge:
    __slots__=('to','rev','cap','cost')
    def __init__(self,to,rev,cap,cost): self.to,self.rev,self.cap,self.cost=to,rev,cap,cost
def add(g,u,v,cap,cost):
    g[u].append(Edge(v,len(g[v]),cap,cost)); g[v].append(Edge(u,len(g[u])-1,0,-cost))

def exact_floor_transport(rows,cols,quotas,den):
    """Exact L1 optimum when quota floors leave a unit-capacity residual b-matching."""
    row_totals=dict(rows); col_totals=dict(cols); rows=sorted(row_totals); cols=sorted(col_totals); base={}; rr={r:row_totals[r] for r in rows}; cc={c:col_totals[c] for c in cols}
    remainders={}
    for r in rows:
        for c in cols:
            n=quotas.get((r,c))
            if n is None: continue
            q,rem=divmod(n,den); base[r,c]=q; remainders[r,c]=rem; rr[r]-=q; cc[c]-=q
    if any(v<0 for v in rr.values()) or any(v<0 for v in cc.values()) or sum(rr.values())!=sum(cc.values()): raise ValueError('Infeasible floor residual margins')
    if any(v>sum((r,c) in quotas for c in cols) for r,v in rr.items()) or any(v>sum((r,c) in quotas for r in rows) for c,v in cc.items()): raise ValueError('Residual margins exceed support')
    n=2+len(rows)+len(cols); s=n-2;t=n-1;g=[[] for _ in range(n)]; ri={r:i for i,r in enumerate(rows)}; ci={c:len(rows)+i for i,c in enumerate(cols)}
    for r in rows:add(g,s,ri[r],rr[r],0)
    for r in rows:
        for c in cols:
            if (r,c) in quotas:add(g,ri[r],ci[c],1,den-remainders[r,c])
    for c in cols:add(g,ci[c],t,cc[c],0)
    need=sum(rr.values()); flow=0; cost=0; pot=[0]*n
    while flow<need:
        inf=None; dist=[inf]*n; prev=[None]*n; dist[s]=0; heap=[(0,s)]
        while heap:
            d,u=heapq.heappop(heap)
            if d!=dist[u]:continue
            for ei,e in enumerate(g[u]):
                if not e.cap:continue
                nd=d+e.cost+pot[u]-pot[e.to]
                if dist[e.to] is None or nd<dist[e.to]:dist[e.to]=nd;prev[e.to]=(u,ei);heapq.heappush(heap,(nd,e.to))
        if dist[t] is None:raise ValueError('Exact integer margins are infeasible on the permitted support')
        for i,d in enumerate(dist):
            if d is not None:pot[i]+=d
        v=t
        while v!=s:
            u,ei=prev[v];e=g[u][ei];e.cap-=1;g[v][e.rev].cap+=1;cost+=e.cost;v=u
        flow+=1
    out=dict(base)
    for r in rows:
        u=ri[r]
        for e in g[u]:
            if e.to in ci.values() and e.cost>=0:
                c=cols[e.to-len(rows)]
                if (r,c) in quotas and e.cap==0:out[r,c]=out.get((r,c),0)+1
    if any(sum(out.get((r,c),0) for c in cols)!=row_totals[r] for r in rows) or any(sum(out.get((r,c),0) for r in rows)!=col_totals[c] for c in cols):raise AssertionError('Transport did not conserve margins')
    return out,flow,cost

def load_local():
    raw=BOOK.read_bytes()
    if len(raw)!=47_036_552 or hashlib.sha256(raw).hexdigest()!=BOOK_HASH:raise ValueError('Pinned workbook integrity failed')
    mapping={r['officialCode']:r for r in json.loads((GEO/'normalized-local-authorities.json').read_text(encoding='utf-8'))['records']}
    wb=openpyxl.load_workbook(BOOK,read_only=True,data_only=True)
    expected=['Cover sheet','Contents','Notes','Related publications','MYEB1','MYEB2','MYEB3','MYEB4','MYEB5']
    if wb.sheetnames!=expected:raise ValueError('Workbook sheet inventory changed')
    ws=wb['MYEB1']; header=[c.value for c in ws[2]]
    if header[:5]!=['ladcode23','laname23','country','sex','age'] or header[-1]!='population_2024':raise ValueError('MYEB1 schema changed')
    areas={}
    for row in ws.iter_rows(min_row=3,values_only=True):
        if row[0] is None:continue
        code,name,country,sex,age=row[:5]; value=row[18]
        if code not in mapping or name!=mapping[code]['displayName'] or country!=code[0] or sex not in ('f','m') or not isinstance(age,int) or not 0<=age<=90 or not isinstance(value,int) or value<0:raise ValueError(f'Malformed MYEB1 row for {code}')
        rec=areas.setdefault(code,{'officialCode':code,'areaId':mapping[code]['placeId'],'displayName':name,'countryCode':country,'cells':{}})
        if (sex,age) in rec['cells']:raise ValueError('Duplicate area/sex/age cell')
        rec['cells'][sex,age]=value
    if len(areas)!=361 or any(len(a['cells'])!=182 for a in areas.values()) or set(areas)!=set(mapping):raise ValueError('Workbook/geography area inventory mismatch')
    out=[]
    for code,a in sorted(areas.items()):
        ages=[a['cells']['f',age]+a['cells']['m',age] for age in range(90)]; old=a['cells']['f',90]+a['cells']['m',90]
        out.append({k:a[k] for k in ('officialCode','areaId','displayName','countryCode')}|{'completedAges':ages,'age90Plus':old,'total':sum(ages)+old})
    return out

def self_test():
    # Exact one-edge, two-column rounding, canonical equal-cost tie and order independence.
    out,_,_=exact_floor_transport({'a':7},{'x':7},{('a','x'):7},1);assert out=={('a','x'):7}
    rows={'area.a':1,'area.b':1};cols={'year.1':1,'year.2':1};q={(r,c):1 for r in rows for c in cols}
    first=exact_floor_transport(rows,cols,q,2)[0];second=exact_floor_transport(dict(reversed(list(rows.items()))),dict(reversed(list(cols.items()))),dict(reversed(list(q.items()))),2)[0]
    assert first==second=={('area.a','year.1'):1,('area.a','year.2'):0,('area.b','year.1'):0,('area.b','year.2'):1}
    try:exact_floor_transport({'a':2},{'x':2},{('a','x'):0},1)
    except ValueError:pass
    else:raise AssertionError('Unsupported-edge infeasibility was accepted')
    # The June-30 compatible interval is 365 days normally and 366 when the later birth year is leap.
    assert 184+(182 if 2024%4==0 else 181)==366 and 184+181==365
    print('UK geographic allocation exact-solver synthetic fixtures passed.')

def main():
    started=time.perf_counter(); areas=load_local();
    if '--reverse-source' in sys.argv:areas=list(reversed(areas))
    v2=json.loads(V2.read_text(encoding='utf-8'))
    targets={m['dimensions']['birthYear']:int(m['value']) for m in v2['measures'] if m['dimensions']['kind']=='birth-year'}
    if len(targets)!=119 or sum(targets.values())!=TOTAL:raise ValueError('Frozen v2 birth-year margins changed')
    area_rows={a['areaId']:a for a in areas}; national_age=[sum(a['completedAges'][c] for a in areas) for c in range(90)]; national_old=sum(a['age90Plus'] for a in areas)
    if sum(a['total'] for a in areas)!=TOTAL:raise ValueError('Local evidence does not conserve UK total')
    # Recover exact national earlier-year margins recursively from frozen v2.
    earlier=[]
    for c in range(90):
        later_year=2024-c; carry=earlier[c-1] if c else 0
        later=targets[later_year]-carry; e=national_age[c]-later
        if not 0<=e<=national_age[c]:raise ValueError(f'National margin conflict at completed age {c}')
        earlier.append(e)
    old_targets={y:targets[y] for y in range(1906,1934)}; old_targets[1934]=targets[1934]-earlier[89]
    if any(v<0 for v in old_targets.values()) or sum(old_targets.values())!=national_old:raise ValueError('Frozen old-age birth years conflict with local 90+ evidence')
    latent=defaultdict(int); residuals=0; transport_cost=0
    for c in range(90):
        ey,ly=2023-c,2024-c; den=366 if (ly%4==0 and (ly%100!=0 or ly%400==0)) else 365
        rows={a['areaId']:a['completedAges'][c] for a in areas}; cols={str(ey):earlier[c],str(ly):national_age[c]-earlier[c]}
        quotas={(aid,str(ey)):n*184 for aid,n in rows.items()}|{(aid,str(ly)):n*(den-184) for aid,n in rows.items()}
        alloc,f,cost=exact_floor_transport(rows,cols,quotas,den);residuals+=f;transport_cost+=cost
        for (aid,y),n in alloc.items():latent[aid,c,int(y)]+=n
    rows={a['areaId']:a['age90Plus'] for a in areas}; cols={str(y):n for y,n in old_targets.items()}; quotas={(aid,y):n*cols[y] for aid,n in rows.items() for y in cols}
    old_alloc,old_residuals,old_cost=exact_floor_transport(rows,cols,quotas,national_old);residuals+=old_residuals;transport_cost+=old_cost
    for (aid,y),n in old_alloc.items():latent[aid,90,int(y)]+=n
    total_deviation=Fraction(0);max_deviation=Fraction(0)
    for a in areas:
        aid=a['areaId']
        for c,n in enumerate(a['completedAges']):
            ey,ly=2023-c,2024-c;den=366 if (ly%4==0 and (ly%100!=0 or ly%400==0)) else 365
            for y,w in ((ey,184),(ly,den-184)):
                deviation=abs(Fraction(latent[aid,c,y])-Fraction(n*w,den));total_deviation+=deviation;max_deviation=max(max_deviation,deviation)
        for y,national in old_targets.items():
            deviation=abs(Fraction(latent[aid,90,y])-Fraction(a['age90Plus']*national,national_old));total_deviation+=deviation;max_deviation=max(max_deviation,deviation)
    counts=defaultdict(int)
    for (aid,c,y),n in latent.items():
        if n<0:raise AssertionError('Negative cell')
        counts[aid,y]+=n
    cohorts=[{'id':cohort_id(a,y),'countryId':'uk','areaId':a,'birthYear':y,'generationProfileId':PROFILE_ID,'count':n} for (a,y),n in counts.items() if n]
    cohorts.sort(key=lambda x:x['id'])
    package_sem={'schemaVersion':1,'id':PACKAGE_ID,'status':'candidate','countryId':'uk','effectiveDate':{'year':2024,'month':6,'day':30},'populationUniverseId':'population.universe.uk-usual-residents-mid-2024-v1','sourcePopulationPackageId':'uk.population.mid-2024.v2','sourcePopulationFingerprint':v2['fingerprint'],'geographyPartitionId':PARTITION_ID,'geographyPartitionFingerprint':'fnv1a64-v1:3d1a3446a16c58cb','sourceArtifactSha256':BOOK_HASH,'generationProfileId':PROFILE_ID,'cohorts':cohorts}
    package=package_sem|{'fingerprint':fnv(package_sem)}
    # Exact audits.
    by_year=defaultdict(int);by_area=defaultdict(int);by_country=defaultdict(int)
    code_to_area={a['areaId']:a for a in areas}
    for c in cohorts:by_year[c['birthYear']]+=c['count'];by_area[c['areaId']]+=c['count'];by_country[code_to_area[c['areaId']]['countryCode']]+=c['count']
    if dict(by_year)!=targets or any(by_area[a['areaId']]!=a['total'] for a in areas):raise AssertionError('Final margins failed')
    country_expected={'E':58_620_101,'W':3_186_581,'S':5_546_900,'N':1_927_855}
    if dict(by_country)!=country_expected:raise AssertionError(f'Country totals differ: {dict(by_country)}')
    pop={'version':1,'coverage':[{'countryId':'uk','status':'complete','source':PACKAGE_ID,'areaPartitionId':PARTITION_ID}],'cohorts':cohorts,'memberships':[]}
    report={'version':1,'candidatePackageId':PACKAGE_ID,'candidateFingerprint':package['fingerprint'],'status':'candidate-awaiting-persistence-gate','sourcePopulationPackageId':'uk.population.mid-2024.v2','sourcePopulationFingerprint':v2['fingerprint'],'sourceArtifactSha256':BOOK_HASH,'geographyPartitionId':PARTITION_ID,'geographyPartitionFingerprint':'fnv1a64-v1:3d1a3446a16c58cb','areaCount':361,'birthYearCount':119,'birthYearRange':[1906,2024],'latentSupportCount':361*90*2+361*len(old_targets),'nonzeroLatentCount':sum(n>0 for n in latent.values()),'nonzeroFinalCohortCount':len(cohorts),'checks':{'localCompletedAgeMargins':361*90,'localCompletedAgeExact':True,'local90PlusMargins':361,'local90PlusExact':True,'nationalBirthYearMargins':119,'nationalBirthYearsExact':True,'localAreaMargins':361,'localAreasExact':True,'nationalTotal':TOTAL,'nationalTotalExact':True,'unsupportedEdges':0,'negativeCells':0},'constituentCountryTotals':{'england':by_country['E'],'wales':by_country['W'],'scotland':by_country['S'],'northernIreland':by_country['N']},'gregorianMethod':'Exact valid-birthday-day overlap: July 1 through December 31 (184 days) versus January 1 through June 30 (181 or 182 days).','oldAgeMethod':'Exact local 90+ row margins allocated across frozen v2 old-age birth-year margins as the disclosed within-total prior.','objective':'Minimum exact L1 deviation from rational quotas; floor residual min-cost flow; canonical area/category/year ordering resolves equal costs deterministically.','residualAssignments':residuals,'oldAgeResidualAssignments':old_residuals,'integerTransportCostNumerator':str(transport_cost),'candidatePopulationSerializedBytesUtf8':len(canonical(pop).encode()),'buildDurationMs':round((time.perf_counter()-started)*1000,3),'evidencePrecision':'Canonical exact integers are calibrated/derived simulation values. Local source estimates and frozen national calibration do not have matching empirical precision.'}
    report['status']='candidate-persistence-gate-failed'
    report.pop('integerTransportCostNumerator',None)
    report['quotaDeviation']={'totalExact':f'{total_deviation.numerator}/{total_deviation.denominator}','maximumExact':f'{max_deviation.numerator}/{max_deviation.denominator}'}
    report['persistenceGate']={'result':'failed','representativeGameSerializedBytesUtf8':9_595_304,'observedSaveMs':555.697,'observedLoadMs':254.268,'browserLocalStorageBudgetBytes':5*1024*1024,'browserStorageSupported':False,'saveLoadContinuationPersonId':'person:2','note':'Timing is an informational observation and is excluded from candidate semantics/fingerprint.'}
    write('normalized-local-age-margins.json',{'version':1,'sourceArtifactSha256':BOOK_HASH,'areas':sorted(areas,key=lambda a:a['areaId'])})
    write('compiled-population-v3-candidate.json',package);write('allocation-report.json',report)
    print(json.dumps({'fingerprint':package['fingerprint'],'cohorts':len(cohorts),'populationBytes':report['candidatePopulationSerializedBytesUtf8'],'durationMs':report['buildDurationMs']},indent=2))

if __name__=='__main__':self_test() if '--self-test' in sys.argv else main()
